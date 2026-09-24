"""Guided one-time import of the 2024 Excel workbook (spec section 8).

    python manage.py import_workbook "Stock 2024.xlsx"             # dry run: validate + reconcile, nothing saved
    python manage.py import_workbook "Stock 2024.xlsx" --commit    # actually import
    python manage.py import_workbook book.xlsx --report import-report.xlsx

Sheets are recognised by name (Stock File / Silver / Palladium / Day Book / Profit / Misc) and columns by
header text, so small layout differences between sheets are tolerated. Every error cell (#DIV/0!, #REF!, ...)
and every row that cannot be imported is listed in the report for manual review.
Import order follows the spec: stock -> customers/suppliers -> sales -> cash book -> investors.
"""
import datetime
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from cashbook.models import CashBookEntry
from cashbook.services import ensure_default_heads
from core import calculations as calc
from core.models import HeadOfAccount, Sequence, ShopSettings, User
from investors.models import Investor, InvestorTransaction
from parties.models import Customer, Supplier
from sales.models import Invoice, Payment, SaleLine
from stock.models import StockItem
from stock.services import bump_sequence_for

ERROR_VALUES = {"#DIV/0!", "#REF!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}

SHEETS = {
    "stock": ["stock file", "stock"],
    "silver": ["silver"],
    "palladium": ["palladium"],
    "daybook": ["day book", "daybook", "cash book"],
    "profit": ["profit"],
    "misc": ["misc"],
}

STOCK_COLS = {
    "code": ["code", "item code", "bin", "bin no", "tag", "tag no", "item no", "sr code"],
    "category": ["category", "item", "item type", "article", "cat"],
    "name": ["name", "description", "item name", "detail", "particulars"],
    "pieces": ["pcs", "pieces", "qty", "quantity", "no of pcs"],
    "type": ["type", "plain/jurao", "plain jurao"],
    "size": ["size", "size inch", "inches"],
    "design": ["design", "design no", "design name"],
    "supplier": ["supplier", "party", "karigar", "from", "supplier name"],
    "nature": ["nature"],
    "source": ["source", "stock type", "opening/current"],
    "book_no": ["book", "book no", "book no."],
    "page_no": ["page", "page no", "page no."],
    "gross": ["gross", "gross wt", "gross weight", "weight", "wt", "total wt"],
    "stone": ["stone", "big stone", "stone wt", "big stone wt"],
    "pd": ["pd", "palladium", "diamond", "pd/diamond", "pd wt", "palladium/diamond"],
    "extra": ["extra", "extra/less", "less", "extra less", "+/-"],
    "ratti": ["ratti", "ratti kaat", "kaat", "purity", "rati"],
    "rate": ["rate", "gold rate", "purchase rate", "p rate"],
    "extra_costs": ["extra cost", "extra costs", "beads", "dropper", "stone price", "diamond price", "silver price",
                    "beads/dropper", "other cost"],
    "purchase_date": ["date", "purchase date", "p date", "pur date"],
    "sale_date": ["sale date", "sold date", "sold on", "s date"],
    "customer": ["customer", "customer name", "sold to", "buyer"],
    "sale_amount": ["sale", "sale price", "sale amount", "sold price", "net sale", "sale rs", "sold amount"],
    "status": ["status"],
    "sale_book": ["sale book", "s book"],
    "sale_page": ["sale page", "s page"],
}

DAYBOOK_COLS = {
    "date": ["date"],
    "type": ["type", "i/e", "ie"],
    "hoa": ["hoa", "head", "head of account", "head of a/c"],
    "sub_hoa": ["sub hoa", "sub head", "sub-head", "subhead"],
    "detail": ["detail", "details"],
    "debited_to": ["debited to", "debit to", "debited"],
    "party": ["party", "received from", "paid to", "received from/paid to", "name"],
    "book_no": ["book", "book no"],
    "page_no": ["page", "page no"],
    "pound": ["pound", "pound wt", "pound weight"],
    "pasa": ["pasa", "pasa wt", "pasa weight"],
    "rate": ["rate"],
    "amount": ["amount", "amount rs", "rs"],
    "income": ["income", "in", "receipt"],
    "expense": ["expense", "expenditure", "out", "payment"],
    "account": ["account", "cash account", "a/c"],
}

PROFIT_COLS = {
    "name": ["name", "investor", "investor name"],
    "reference": ["reference", "ref"],
    "book_no": ["book", "book no"],
    "page_no": ["page", "page no"],
    "share": ["share", "share %", "ratio", "profit ratio", "%"],
    "amount": ["amount", "investment", "capital", "base value"],
    "date": ["date"],
}

MISC_COLS = {
    "date": ["date"],
    "description": ["description", "item", "detail", "particulars"],
    "customer": ["customer", "name"],
    "amount": ["amount", "sale", "price", "rs"],
    "cost": ["cost", "purchase"],
}


def norm(text):
    return re.sub(r"[^a-z0-9/%+\-]+", " ", str(text or "").lower()).strip()


def clean_name(text):
    return re.sub(r"\s+", " ", str(text or "")).strip().title()


class Issues:
    def __init__(self):
        self.rows = []

    def add(self, sheet, row, column, problem, value=""):
        self.rows.append({"sheet": sheet, "row": row, "column": column, "problem": problem, "value": str(value)[:80]})


class Command(BaseCommand):
    help = "Import the 2024 workbook (dry run unless --commit)."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--commit", action="store_true", help="Save the import (default is a dry run)")
        parser.add_argument("--sheets", default="stock,silver,palladium,daybook,profit,misc")
        parser.add_argument("--report", help="Write issues + reconciliation to this .xlsx file")
        parser.add_argument("--user", default=None, help="Username recorded as creator (default: first owner)")

    # ---------- helpers ----------
    def _find_sheet(self, wb, key):
        for ws in wb.worksheets:
            n = norm(ws.title)
            if any(alias == n or n.startswith(alias) for alias in SHEETS[key]):
                return ws
        return None

    def _header(self, ws, spec):
        best = (0, None, {})
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            mapping = {}
            for c_idx, cell in enumerate(row):
                h = norm(cell)
                if not h:
                    continue
                for field, aliases in spec.items():
                    if field in mapping and field != "extra_costs":
                        continue
                    if h in aliases:
                        if field == "extra_costs":
                            mapping.setdefault(field, []).append(c_idx)
                        else:
                            mapping[field] = c_idx
                        break
            if len(mapping) > best[0]:
                best = (len(mapping), r_idx, mapping)
        return best[1], best[2]

    def _val(self, row, mapping, field, sheet, r, issues):
        idx = mapping.get(field)
        if idx is None or idx >= len(row):
            return None
        v = row[idx]
        if isinstance(v, str) and v.strip() in ERROR_VALUES:
            issues.add(sheet, r, field, "Error cell in workbook", v)
            return None
        return v

    def _dec(self, v, sheet, r, field, issues, default=None):
        if v is None or (isinstance(v, str) and not v.strip()):
            return default
        try:
            return Decimal(str(v).replace(",", "").strip())
        except InvalidOperation:
            issues.add(sheet, r, field, "Not a number", v)
            return default

    def _date(self, v, sheet, r, field, issues):
        if v is None or v == "":
            return None
        if isinstance(v, datetime.datetime):
            return v.date()
        if isinstance(v, datetime.date):
            return v
        if isinstance(v, (int, float)):
            try:
                return from_excel(v).date()
            except Exception:  # noqa: BLE001
                issues.add(sheet, r, field, "Bad Excel date serial", v)
                return None
        s = str(v).strip()
        for fmt in ("%d/%m/%y", "%d/%m/%Y", "%d-%m-%Y", "%d-%m-%y", "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        issues.add(sheet, r, field, "Unrecognised date", v)
        return None

    # ---------- main ----------
    def handle(self, *args, **opts):
        try:
            wb = load_workbook(opts["path"], data_only=True, read_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {opts['path']}")
        self.user = (User.objects.filter(username=opts["user"]).first() if opts["user"]
                     else User.objects.filter(role=User.Role.OWNER).first())
        wanted = [s.strip() for s in opts["sheets"].split(",") if s.strip()]
        issues = Issues()
        stats = defaultdict(lambda: defaultdict(Decimal))
        self.customers, self.suppliers = {}, {}

        with transaction.atomic():
            ensure_default_heads()
            for key, metal in (("stock", "gold"), ("silver", "silver"), ("palladium", "palladium")):
                if key in wanted:
                    ws = self._find_sheet(wb, key)
                    if ws:
                        self.import_stock(ws, metal, issues, stats[ws.title])
                    else:
                        issues.add(key, "-", "-", "Sheet not found")
            for key, fn in (("misc", self.import_misc), ("daybook", self.import_daybook),
                            ("profit", self.import_profit)):
                if key in wanted:
                    ws = self._find_sheet(wb, key)
                    if ws:
                        fn(ws, issues, stats[ws.title])
                    else:
                        issues.add(key, "-", "-", "Sheet not found")
            for c in Customer.objects.all():
                c.recompute_balance()
            for s in Supplier.objects.all():
                s.recompute_balance()
            if not opts["commit"]:
                transaction.set_rollback(True)

        self._print(stats, issues, opts["commit"])
        if opts.get("report"):
            self._write_report(opts["report"], stats, issues)

    def _customer(self, name):
        name = clean_name(name)
        if not name:
            return None
        key = name.lower()
        if key not in self.customers:
            c = Customer.objects.filter(name__iexact=name).first() or Customer.objects.create(name=name)
            self.customers[key] = c
        return self.customers[key]

    def _supplier(self, name):
        name = clean_name(name)
        if not name:
            return None
        key = name.lower()
        if key not in self.suppliers:
            s = Supplier.objects.filter(name__iexact=name).first() or Supplier.objects.create(name=name)
            self.suppliers[key] = s
        return self.suppliers[key]

    def _choice(self, value, choices, default):
        v = norm(value)
        for c in choices:
            if v == c or v.startswith(c[:3]):
                return c
        return default

    def import_stock(self, ws, metal, issues, st):
        header_row, m = self._header(ws, STOCK_COLS)
        sheet = ws.title
        if not header_row or "gross" not in m:
            issues.add(sheet, "-", "-", "Could not find a header row with a weight column")
            return
        shop = ShopSettings.load()
        prefixes = dict(shop.code_prefixes)
        for r, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(v not in (None, "") for v in row):
                continue
            g = lambda f: self._val(row, m, f, sheet, r, issues)  # noqa: E731
            gross = self._dec(g("gross"), sheet, r, "gross", issues)
            if gross is None or gross <= 0:
                issues.add(sheet, r, "gross", "Missing/zero gross weight - row skipped", g("gross"))
                continue
            code = str(g("code") or "").strip().upper().replace(" ", "")
            category = clean_name(g("category")) or "Misc"
            if not code:
                issues.add(sheet, r, "code", "Missing item code - row skipped")
                continue
            if StockItem.objects.filter(code=code).exists():
                issues.add(sheet, r, "code", "Duplicate code - row skipped", code)
                continue
            letters = re.match(r"^[A-Z]+", code)
            if letters and category not in prefixes and letters.group(0) not in prefixes.values():
                prefixes[category] = letters.group(0)
            extra_costs = Decimal("0")
            for idx in m.get("extra_costs", []):
                v = row[idx] if idx < len(row) else None
                if isinstance(v, str) and v.strip() in ERROR_VALUES:
                    issues.add(sheet, r, "extra_costs", "Error cell in workbook", v)
                    continue
                extra_costs += self._dec(v, sheet, r, "extra_costs", issues, Decimal("0"))
            ratti = self._dec(g("ratti"), sheet, r, "ratti", issues, Decimal("96") if metal != "gold" else None)
            if ratti is None:
                issues.add(sheet, r, "ratti", "Missing ratti kaat - assumed 90")
                ratti = Decimal("90")
            nature = self._choice(g("nature"), ["new", "transfer", "karigar", "old"], "new")
            source = "opening_stock" if "open" in norm(g("source")) else "current_purchase"
            sale_date = self._date(g("sale_date"), sheet, r, "sale_date", issues)
            sale_amount = self._dec(g("sale_amount"), sheet, r, "sale_amount", issues)
            status_text = norm(g("status"))
            sold = bool(sale_date or sale_amount) or status_text.startswith("sold")
            item = StockItem.objects.create(
                code=code, metal=metal, category=category, name=str(g("name") or "")[:200],
                pieces=int(self._dec(g("pieces"), sheet, r, "pieces", issues, Decimal(1))),
                type=self._choice(g("type"), ["plain", "jurao"], "plain"), size=str(g("size") or "")[:30],
                design=str(g("design") or "")[:100], supplier=self._supplier(g("supplier")), nature=nature,
                source_type=source, book_no=str(g("book_no") or "")[:30], page_no=str(g("page_no") or "")[:20],
                gross_weight=gross,
                big_stone_weight=self._dec(g("stone"), sheet, r, "stone", issues, Decimal(0)),
                pd_diamond_weight=self._dec(g("pd"), sheet, r, "pd", issues, Decimal(0)),
                extra_less_gold=self._dec(g("extra"), sheet, r, "extra", issues, Decimal(0)),
                ratti_kaat=ratti, purchase_rate=self._dec(g("rate"), sheet, r, "rate", issues, Decimal(0)),
                extra_costs=extra_costs,
                purchase_date=self._date(g("purchase_date"), sheet, r, "purchase_date", issues),
                status=StockItem.Status.SOLD if sold else StockItem.Status.IN_STOCK, sold_date=sale_date,
                notes="Imported from 2024 workbook", created_by=self.user, updated_by=self.user,
            )
            bump_sequence_for(code)
            st["items"] += 1
            st["gross_weight"] += item.gross_weight
            st["pasa"] += item.pasa
            st["total_cost"] += item.total_cost
            if sold:
                if not sale_amount:
                    issues.add(sheet, r, "sale_amount", "Sold item without sale amount - invoice not created", code)
                    continue
                self._legacy_invoice(
                    date=sale_date or item.purchase_date or datetime.date(2024, 1, 1),
                    customer=self._customer(g("customer")), book=str(g("sale_book") or g("book_no") or ""),
                    page=str(g("sale_page") or g("page_no") or ""),
                    lines=[dict(stock_item=item, description=f"{item.code} {item.category}", weight=item.gold_weight,
                                ratti_kaat=item.ratti_kaat, pasa=item.pasa, gold_value=sale_amount,
                                line_total=sale_amount, purchase_cost=item.total_cost,
                                profit=sale_amount - item.total_cost)],
                )
                st["sales"] += sale_amount
                st["sold_items"] += 1
        shop.code_prefixes = prefixes
        shop.save()

    def _legacy_invoice(self, date, customer, book, page, lines):
        total = sum((l["line_total"] for l in lines), Decimal("0"))
        number = f"OLD-{Sequence.next_value('legacy_invoice'):06d}"
        inv = Invoice.objects.create(
            number=number, date=date, customer=customer, legacy_book=book[:30], legacy_page=page[:20],
            sale_rate=Decimal("0"), subtotal=total, net_amount=total, salesperson=self.user,
            notes="Imported from 2024 workbook (assumed fully paid)",
        )
        for l in lines:
            SaleLine.objects.create(invoice=inv, **l)
        # Historical sales are assumed paid; their cash is already in the imported Day Book, so no cash-book entry.
        Payment.objects.create(invoice=inv, customer=customer, kind=Payment.Kind.SALE, method=Payment.Method.CASH,
                               amount=total, date=date, notes="Imported", created_by=self.user)
        inv.refresh_totals()
        return inv

    def import_misc(self, ws, issues, st):
        header_row, m = self._header(ws, MISC_COLS)
        sheet = ws.title
        if not header_row or "amount" not in m:
            issues.add(sheet, "-", "-", "Could not find header with an amount column")
            return
        for r, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(v not in (None, "") for v in row):
                continue
            g = lambda f: self._val(row, m, f, sheet, r, issues)  # noqa: E731
            amount = self._dec(g("amount"), sheet, r, "amount", issues)
            if not amount:
                continue
            cost = self._dec(g("cost"), sheet, r, "cost", issues, Decimal(0))
            desc = str(g("description") or "Misc sale")[:200]
            self._legacy_invoice(
                date=self._date(g("date"), sheet, r, "date", issues) or datetime.date(2024, 1, 1),
                customer=self._customer(g("customer")), book="", page="",
                lines=[dict(description=desc, gold_value=Decimal(0), line_total=amount, stone_value=Decimal(0),
                            making_charges=amount, purchase_cost=cost, profit=amount - cost)],
            )
            st["invoices"] += 1
            st["sales"] += amount

    def import_daybook(self, ws, issues, st):
        header_row, m = self._header(ws, DAYBOOK_COLS)
        sheet = ws.title
        if not header_row or "date" not in m or not ({"amount", "income", "expense"} & set(m)):
            issues.add(sheet, "-", "-", "Could not find header with date and amount columns")
            return
        heads = {}
        accounts = ShopSettings.load().cash_accounts

        def head(name, type_, parent=None):
            key = (name.lower(), type_, parent.pk if parent else None)
            if key not in heads:
                heads[key] = (HeadOfAccount.objects.filter(name__iexact=name, type=type_, parent=parent).first()
                              or HeadOfAccount.objects.create(name=name, type=type_, parent=parent,
                                                              report_group=parent.report_group if parent else ""))
            return heads[key]

        for r, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(v not in (None, "") for v in row):
                continue
            g = lambda f: self._val(row, m, f, sheet, r, issues)  # noqa: E731
            date = self._date(g("date"), sheet, r, "date", issues)
            if not date:
                issues.add(sheet, r, "date", "Missing date - row skipped")
                continue
            type_ = str(g("type") or "").strip().upper()[:1]
            amount = self._dec(g("amount"), sheet, r, "amount", issues)
            if amount is None:
                inc = self._dec(g("income"), sheet, r, "income", issues)
                exp = self._dec(g("expense"), sheet, r, "expense", issues)
                if inc:
                    type_, amount = "I", inc
                elif exp:
                    type_, amount = "E", exp
            if not amount:
                issues.add(sheet, r, "amount", "Missing amount - row skipped")
                continue
            if amount < 0:
                amount = -amount
                type_ = "E" if type_ == "I" else "I"
            if type_ not in ("I", "E"):
                issues.add(sheet, r, "type", "Missing I/E type - row skipped", g("type"))
                continue
            hoa_name = clean_name(g("hoa")) or "Unclassified"
            hoa = head(hoa_name, type_)
            sub_name = clean_name(g("sub_hoa"))
            sub = head(sub_name, type_, hoa) if sub_name else None
            account = str(g("account") or "").strip()
            detail = str(g("detail") or "").strip()
            if account not in accounts:
                account = detail if detail in accounts else "shop_cash"
            CashBookEntry.objects.create(
                date=date, type=type_, hoa=hoa, sub_hoa=sub, detail=detail[:100],
                debited_to=str(g("debited_to") or "")[:100], party=str(g("party") or "")[:200],
                book_no=str(g("book_no") or "")[:30], page_no=str(g("page_no") or "")[:20],
                pound_weight=self._dec(g("pound"), sheet, r, "pound", issues),
                pasa_weight=self._dec(g("pasa"), sheet, r, "pasa", issues),
                rate=self._dec(g("rate"), sheet, r, "rate", issues), amount=calc.q_money(amount) or amount,
                source_type=CashBookEntry.Source.IMPORT, cash_account=account, created_by=self.user,
                updated_by=self.user,
            )
            st["entries"] += 1
            st["income" if type_ == "I" else "expenditure"] += amount

    def import_profit(self, ws, issues, st):
        header_row, m = self._header(ws, PROFIT_COLS)
        sheet = ws.title
        if not header_row or "name" not in m:
            issues.add(sheet, "-", "-", "Could not find header with an investor name column")
            return
        for r, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            g = lambda f: self._val(row, m, f, sheet, r, issues)  # noqa: E731
            name = clean_name(g("name"))
            if not name:
                continue
            share = self._dec(g("share"), sheet, r, "share", issues, Decimal(0))
            if share and share < 1:  # 0.035 written as a fraction
                share = share * 100
            inv, created = Investor.objects.get_or_create(name=name, defaults={
                "reference": str(g("reference") or "")[:100], "book_no": str(g("book_no") or "")[:30],
                "page_no": str(g("page_no") or "")[:20], "share_percent": share,
                "notes": "Imported from 2024 Profit sheet - recalculate periods in system"})
            st["investors"] += 1 if created else 0
            amount = self._dec(g("amount"), sheet, r, "amount", issues)
            if amount:
                InvestorTransaction.objects.create(
                    investor=inv, date=self._date(g("date"), sheet, r, "date", issues) or datetime.date(2024, 1, 1),
                    type="investment", amount=amount, notes="Imported (cash already in Day Book)",
                    created_by=self.user)
                st["investment"] += amount

    def _print(self, stats, issues, committed):
        self.stdout.write(self.style.MIGRATE_HEADING("Reconciliation (compare with the workbook totals):"))
        for sheet, values in stats.items():
            parts = ", ".join(f"{k}={v:,.3f}".rstrip("0").rstrip(".") for k, v in values.items())
            self.stdout.write(f"  {sheet}: {parts}")
        self.stdout.write(self.style.MIGRATE_HEADING(f"Issues to review: {len(issues.rows)}"))
        for i in issues.rows[:60]:
            self.stdout.write(f"  [{i['sheet']} r{i['row']} {i['column']}] {i['problem']} {i['value']}")
        if len(issues.rows) > 60:
            self.stdout.write(f"  ... {len(issues.rows) - 60} more (use --report to get them all)")
        if committed:
            self.stdout.write(self.style.SUCCESS("Import committed."))
        else:
            self.stdout.write(self.style.WARNING("Dry run only - nothing was saved. Re-run with --commit."))

    def _write_report(self, path, stats, issues):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Issues"
        ws.append(["Sheet", "Row", "Column", "Problem", "Value"])
        for i in issues.rows:
            ws.append([i["sheet"], str(i["row"]), i["column"], i["problem"], i["value"]])
        rec = wb.create_sheet("Reconciliation")
        rec.append(["Sheet", "Measure", "Imported total"])
        for sheet, values in stats.items():
            for k, v in values.items():
                rec.append([sheet, k, float(v)])
        wb.save(path)
        self.stdout.write(f"Report written to {path}")
