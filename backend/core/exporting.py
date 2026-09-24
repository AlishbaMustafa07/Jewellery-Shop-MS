"""Tabular exports: Excel (xlsx), CSV and printable HTML (browser 'Save as PDF')."""
import csv
import datetime
import decimal
import io

from django.http import HttpResponse
from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ShopSettings


def _cell(value):
    if isinstance(value, decimal.Decimal):
        return float(value)
    if hasattr(value, "pk"):
        return str(value)
    return value


def _add_sheet(wb, title, columns, rows, totals=None, subtitle=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    r = 1
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=14)
    if subtitle:
        r += 1
        ws.cell(row=r, column=1, value=subtitle).font = Font(italic=True)
    r += 2
    header_fill = PatternFill("solid", fgColor="F3E7C9")
    for c, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=r, column=c, value=label)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    header_row = r
    for row in rows:
        r += 1
        for c, (key, _) in enumerate(columns, start=1):
            value = _cell(row.get(key) if isinstance(row, dict) else getattr(row, key, None))
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, (datetime.date, datetime.datetime)):
                cell.number_format = "DD/MM/YYYY"
            elif isinstance(value, float):
                cell.number_format = "#,##0.00"
    if totals:
        r += 1
        for c, (key, _) in enumerate(columns, start=1):
            if key in totals:
                cell = ws.cell(row=r, column=c, value=_cell(totals[key]))
                cell.font = Font(bold=True)
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        ws.cell(row=r, column=1).value = ws.cell(row=r, column=1).value or "Total"
        ws.cell(row=r, column=1).font = Font(bold=True)
    for c in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    return ws


def xlsx_response(filename, sheets):
    """sheets: list of dicts {title, columns, rows, totals?, subtitle?}."""
    wb = Workbook()
    for i, s in enumerate(sheets):
        _add_sheet(wb, s["title"], s["columns"], s["rows"], s.get("totals"), s.get("subtitle"), first=(i == 0))
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return resp


def csv_response(filename, columns, rows):
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    resp.write("﻿")  # Excel-friendly BOM
    writer = csv.writer(resp)
    writer.writerow([label for _, label in columns])
    for row in rows:
        out = []
        for key, _ in columns:
            v = row.get(key) if isinstance(row, dict) else getattr(row, key, None)
            if isinstance(v, (datetime.date,)):
                v = v.strftime("%d/%m/%Y")
            out.append(_cell(v))
        writer.writerow(out)
    return resp


def html_response(title, columns, rows, totals=None, subtitle=None, extra=None):
    html = render_to_string("print/table.html", {
        "title": title, "subtitle": subtitle, "columns": columns,
        "rows": [[(row.get(k) if isinstance(row, dict) else getattr(row, k, None)) for k, _ in columns] for row in rows],
        "totals": [totals.get(k, "") if totals else "" for k, _ in columns] if totals else None,
        "shop": ShopSettings.load(), "extra": extra or [],
    })
    return HttpResponse(html)


def export(fmt, filename, title, columns, rows, totals=None, subtitle=None):
    """Return an HttpResponse for fmt in xlsx|csv|html|pdf, or None for JSON."""
    fmt = (fmt or "").lower()
    if fmt == "xlsx":
        return xlsx_response(filename, [dict(title=title, columns=columns, rows=rows, totals=totals, subtitle=subtitle)])
    if fmt == "csv":
        return csv_response(filename, columns, rows)
    if fmt in ("html", "pdf", "print"):
        return html_response(title, columns, rows, totals, subtitle)
    return None
